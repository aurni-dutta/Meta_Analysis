from pathlib import Path
from openpyxl import load_workbook

"""Food-demand scenario calculations."""

def predict_quantity(baseline, elasticity, shock):
    """
    Estimate quantity after a price or food-spending shock.

    Args:
        baseline: Starting food supply quantity.
        elasticity: Relevant price or spending elasticity.
        shock: Fractional change, e.g. 0.10 for +10%.

    Returns:
        Predicted quantity in the same units as baseline.
    """
    return baseline * (1 + elasticity * shock)


from pathlib import Path
from openpyxl import load_workbook



def load_food_group_data(workbook_path, food_group):
    """
    Read baseline food supply and elasticities from Andrew's workbook.
    """
    workbook_path = Path(workbook_path)

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True
    )

    try:
        elasticity_sheet = workbook["TABLE 1 ELASTICITY"]
        consumption_sheet = workbook["TABLE 2 CONSUMPTION"]

        own_price = None
        income = None
        baseline = None

        # Find the requested food group in the elasticity table.
        for row in elasticity_sheet.iter_rows(
            min_row=5,
            values_only=True
        ):
            if row[0] == food_group:
                income = row[11]       # Column L
                own_price = row[12]    # Column M
                break

        # Find the baseline food supply in the consumption table.
        for row in consumption_sheet.iter_rows(
            min_row=3,
            values_only=True
        ):
            if row[0] == food_group:
                baseline = row[2]      # Column C, g/person/day
                break

        if not isinstance(baseline, (int, float)):
            raise ValueError(
                f"Baseline unavailable: {food_group}"
            )

        # Missing estimates are represented by None, not zero.
        if not isinstance(own_price, (int, float)):
            own_price = None

        if not isinstance(income, (int, float)):
            income = None

        return {
            "food_group": food_group,
            "baseline_g_day": baseline,
            "own_price_elasticity": own_price,
            "income_elasticity": income
        }

    finally:
        workbook.close()


def calculate_nutrient_changes(workbook_path, scenario, shock):
    """
    Calculate national food-supply nutrient changes.

    Uses the proportional nutrient-supply method documented
    in Andrew's workbook.

    shock: fractional change, e.g. 0.10 for +10%.
    """
    if scenario not in ("own", "income"):
        raise ValueError("Unsupported nutrient scenario")

    if isinstance(shock, bool) or not isinstance(shock, (int, float)):
        raise ValueError("Shock must be numeric")

    if not -0.30 <= shock <= 0.30:
        raise ValueError("Shock must be between -0.30 and 0.30")

    nutrient_columns = {
        "Energy": 3,
        "Protein": 4,
        "Fat": 5,
        "Iron": 12,
        "Zinc": 13,
        "Vitamin A": 14,
    }

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True
    )

    try:
        elasticity_sheet = workbook["TABLE 1 ELASTICITY"]
        consumption_sheet = workbook["TABLE 2 CONSUMPTION"]

        elasticities = {}

        for row in elasticity_sheet.iter_rows(
            min_row=5,
            values_only=True
        ):
            group = row[0]
            column = 12 if scenario == "own" else 11
            elasticity = row[column]

            if isinstance(group, str) and isinstance(
                elasticity, (int, float)
            ):
                elasticities[group] = elasticity

        changes = {
            nutrient: 0.0
            for nutrient in nutrient_columns
        }

        totals = {}
        missing_groups = []

        for row in consumption_sheet.iter_rows(
            min_row=3,
            values_only=True
        ):
            group = row[0]

            if group == "Total (all groups)":
                totals = {
                    nutrient: row[index]
                    for nutrient, index in nutrient_columns.items()
                }
                continue

            if not isinstance(group, str):
                continue

            if group not in elasticities:
                if group not in (
                    "FAO Grand Total ",
                    "Tobacco"
                ) and isinstance(row[2], (int, float)) and row[2] > 0:
                    missing_groups.append(group)
                continue

            quantity_change = elasticities[group] * shock

            for nutrient, index in nutrient_columns.items():
                baseline = row[index]

                if isinstance(baseline, (int, float)):
                    changes[nutrient] += baseline * quantity_change

        results = {}

        for nutrient in nutrient_columns:
            total = totals.get(nutrient)

            if not isinstance(total, (int, float)) or total <= 0:
                raise ValueError(
                    f"Invalid national baseline: {nutrient}"
                )

            results[nutrient] = round(
                changes[nutrient] / total * 100,
                4
            )

        return {
            "scenario": scenario,
            "shock": shock,
            "nutrient_changes_percent": results,
            "missing_elasticity_groups": missing_groups,
        }

    finally:
        workbook.close()


def calculate_targeted_price_shock(workbook_path, food_group, shock):
    """
    Simulate an own-price shock affecting one food group only.

    Assumptions:
    - Only the selected group's own price changes.
    - Other groups remain at their baseline supply levels.
    - Cross-price substitution is not modeled.
    - Nutrient results are national food-supply changes,
      not changes in individual dietary intake.
    """
    import math

    if not isinstance(food_group, str) or not food_group.strip():
        raise ValueError("Food group is required")

    if (
        isinstance(shock, bool)
        or not isinstance(shock, (int, float))
        or not math.isfinite(shock)
        or not -0.30 <= shock <= 0.30
    ):
        raise ValueError("Shock must be between -0.30 and 0.30")

    nutrient_columns = {
        "Energy": 3,
        "Protein": 4,
        "Fat": 5,
        "Iron": 12,
        "Zinc": 13,
        "Vitamin A": 14,
    }

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        elasticity_sheet = workbook["TABLE 1 ELASTICITY"]
        consumption_sheet = workbook["TABLE 2 CONSUMPTION"]

        elasticity = None
        elasticity_found = False

        for row in elasticity_sheet.iter_rows(
            min_row=5, values_only=True
        ):
            if row[0] == food_group:
                elasticity_found = True
                raw = row[12]  # Own-price elasticity, column M

                if (
                    isinstance(raw, (int, float))
                    and not isinstance(raw, bool)
                    and math.isfinite(raw)
                ):
                    elasticity = raw

                break

        baseline_quantity = None
        group_nutrients = None
        national_totals = {}

        for row in consumption_sheet.iter_rows(
            min_row=3, values_only=True
        ):
            group = row[0]

            if group == "Total (all groups)":
                national_totals = {
                    nutrient: row[index]
                    for nutrient, index in nutrient_columns.items()
                }

            elif group == food_group:
                baseline_quantity = row[2]
                group_nutrients = {
                    nutrient: row[index]
                    for nutrient, index in nutrient_columns.items()
                }

        if (
            not isinstance(baseline_quantity, (int, float))
            or isinstance(baseline_quantity, bool)
            or not math.isfinite(baseline_quantity)
            or baseline_quantity < 0
            or group_nutrients is None
        ):
            raise ValueError("Selected food group's baseline is unavailable")

        if not elasticity_found:
            raise ValueError("Food group not found in elasticity table")

        if elasticity is None:
            raise ValueError(
                "Own-price elasticity is unavailable for this food group"
            )

        quantity_change_fraction = elasticity * shock

        predicted_quantity = baseline_quantity * (
            1 + quantity_change_fraction
        )

        if predicted_quantity < 0:
            raise ValueError("Predicted food supply cannot be negative")

        nutrient_changes = {}

        for nutrient, group_value in group_nutrients.items():
            national_value = national_totals.get(nutrient)

            if (
                not isinstance(national_value, (int, float))
                or isinstance(national_value, bool)
                or not math.isfinite(national_value)
                or national_value <= 0
            ):
                raise ValueError(
                    f"Invalid national baseline: {nutrient}"
                )

            if (
                not isinstance(group_value, (int, float))
                or isinstance(group_value, bool)
                or not math.isfinite(group_value)
            ):
                nutrient_changes[nutrient] = None
                continue

            nutrient_changes[nutrient] = round(
                group_value
                * quantity_change_fraction
                / national_value
                * 100,
                4,
            )

        return {
            "scenario": "targeted_own_price",
            "food_group": food_group,
            "shock": shock,
            "own_price_elasticity": elasticity,
            "baseline_g_day": round(baseline_quantity, 4),
            "predicted_g_day": round(predicted_quantity, 4),
            "quantity_change_percent": round(
                quantity_change_fraction * 100, 4
            ),
            "nutrient_changes_percent": nutrient_changes,
            "cross_price_effects_included": False,
        }

    finally:
        workbook.close()
