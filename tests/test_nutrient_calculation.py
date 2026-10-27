
"""Compare our nutrient calculation with Andrew's Nigeria workbook."""

from pathlib import Path
from openpyxl import load_workbook

WORKBOOK = Path(
    r"C:\Users\aurni\Documents\Work\Workbooks\NIGERIA_FOOD_RESULTS.xlsx"
)

NUTRIENT_COLUMNS = {
    "Energy": 3,       # Excel D
    "Protein": 4,      # Excel E
    "Fat": 5,          # Excel F
    "Iron": 12,        # Excel M
    "Zinc": 13,        # Excel N
    "Vitamin A": 14,   # Excel O
}

SHOCK = 0.10

workbook = load_workbook(
    WORKBOOK,
    read_only=True,
    data_only=True
)

try:
    elasticity_sheet = workbook["TABLE 1 ELASTICITY"]
    consumption_sheet = workbook["TABLE 2 CONSUMPTION"]

    elasticities = {}

    for row in elasticity_sheet.iter_rows(
        min_row=5, values_only=True
    ):
        group = row[0]
        elasticity = row[12]  # Excel M

        if isinstance(group, str) and isinstance(
            elasticity, (int, float)
        ):
            elasticities[group] = elasticity

    totals = {name: 0.0 for name in NUTRIENT_COLUMNS}
    changes = {name: 0.0 for name in NUTRIENT_COLUMNS}

    for row in consumption_sheet.iter_rows(
        min_row=3, values_only=True
    ):
        group = row[0]

        # Exclude totals, notes, and groups without
        # an available own-price elasticity.
        if group not in elasticities:
            continue

        quantity_change = elasticities[group] * SHOCK

        for nutrient, column in NUTRIENT_COLUMNS.items():
            baseline = row[column]

            if isinstance(baseline, (int, float)):
                changes[nutrient] += baseline * quantity_change

    # Read the national baseline totals from Table 2.
    for row in consumption_sheet.iter_rows(
        min_row=3, values_only=True
    ):
        if row[0] == "Total (all groups)":
            for nutrient, column in NUTRIENT_COLUMNS.items():
                totals[nutrient] = row[column]
            break

    print("National nutrient changes:")
    results = {}

    for nutrient in NUTRIENT_COLUMNS:
        baseline_total = totals[nutrient]
        if not isinstance(baseline_total, (int, float)):
            raise ValueError(f"Missing total for {nutrient}")

        percent_change = (
            changes[nutrient] / baseline_total * 100
        )
        results[nutrient] = percent_change

        print(f"{nutrient}: {percent_change:+.3f}%")

    # Compare with TABLE 3's all-groups-at-once row.
    assert abs(results["Energy"] - (-6.114337210385047)) < 0.01
    assert abs(results["Protein"] - (-6.117083990881993)) < 0.01
    assert abs(results["Fat"] - (-7.399723862729451)) < 0.01

    print("\nPASS: Nutrient calculations match workbook totals")

finally:
    workbook.close()
